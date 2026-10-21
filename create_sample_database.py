"""Membuat data/worksheet_database.xlsx (sheet Letters + Templates, contoh lengkap A-T).
JANGAN dijalankan lagi setelah database Anda terisi: file akan ditimpa."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# ---- kata kunci per huruf (CONTOH; ganti sesuai gambar PNG Anda). Image kosong = nama file otomatis dari Word_EN.
LETTERS = [
    ("A", "Apple", "Apel"), ("B", "Ball", "Bola"), ("C", "Cat", "Kucing"), ("D", "Dog", "Anjing"),
    ("E", "Egg", "Telur"), ("F", "Fish", "Ikan"), ("G", "Goat", "Kambing"), ("H", "Hat", "Topi"),
    ("I", "Ice Cream", "Es Krim"), ("J", "Juice", "Jus"), ("K", "Kite", "Layang-Layang"), ("L", "Lion", "Singa"),
    ("M", "Moon", "Bulan"), ("N", "Nest", "Sarang"), ("O", "Orange", "Jeruk"), ("P", "Pen", "Pulpen"),
    ("Q", "Queen", "Ratu"), ("R", "Rabbit", "Kelinci"), ("S", "Sun", "Matahari"), ("T", "Tree", "Pohon"),
    ("U", "Umbrella", "Payung"), ("V", "Violin", "Biola"), ("W", "Whale", "Paus"), ("X", "Xylophone", "Xilofon"),
    ("Y", "Yo-Yo", "Yoyo"), ("Z", "Zebra", "Zebra"),
]
# ---- 20 template A-T
TEMPLATES = [
    ("A", "capital", "A-C", "Capital Letters A\u2013C", "Huruf Besar A\u2013C"),
    ("B", "capital", "D-F", "Capital Letters D\u2013F", "Huruf Besar D\u2013F"),
    ("C", "capital", "G-I", "Capital Letters G\u2013I", "Huruf Besar G\u2013I"),
    ("D", "capital", "J-L", "Capital Letters J\u2013L", "Huruf Besar J\u2013L"),
    ("E", "capital", "M-O", "Capital Letters M\u2013O", "Huruf Besar M\u2013O"),
    ("F", "capital", "P-R", "Capital Letters P\u2013R", "Huruf Besar P\u2013R"),
    ("G", "capital", "S-U", "Capital Letters S\u2013U", "Huruf Besar S\u2013U"),
    ("H", "capital", "V-X", "Capital Letters V\u2013X", "Huruf Besar V\u2013X"),
    ("I", "capital", "Y-Z", "Capital Letters Y\u2013Z", "Huruf Besar Y\u2013Z"),
    ("J", "review", "A-Z", "Capital Letters Review A\u2013Z & Alphabet Order", "Review Huruf Besar A\u2013Z & Urutan Abjad"),
    ("K", "small", "a-c", "Small Letters a\u2013c", "Huruf Kecil a\u2013c"),
    ("L", "small", "d-f", "Small Letters d\u2013f", "Huruf Kecil d\u2013f"),
    ("M", "small", "g-i", "Small Letters g\u2013i", "Huruf Kecil g\u2013i"),
    ("N", "small", "j-l", "Small Letters j\u2013l", "Huruf Kecil j\u2013l"),
    ("O", "small", "m-o", "Small Letters m\u2013o", "Huruf Kecil m\u2013o"),
    ("P", "small", "p-r", "Small Letters p\u2013r", "Huruf Kecil p\u2013r"),
    ("Q", "small", "s-u", "Small Letters s\u2013u", "Huruf Kecil s\u2013u"),
    ("R", "small", "v-z", "Small Letters v\u2013z", "Huruf Kecil v\u2013z"),
    ("S", "similar", "b+d, p+q, m+n", "Telling Similar Letters Apart (b\u2013d, p\u2013q, m\u2013n)", "Membedakan Huruf Mirip (b\u2013d, p\u2013q, m\u2013n)"),
    ("T", "pairs", "A-Z", "Capital & Small Letter Pairs (Aa\u2013Zz)", "Pasangan Huruf Besar\u2013Kecil (Aa\u2013Zz)"),
]

def style(ws, widths, color):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=color)
    for col, w in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

wb = Workbook()
ws = wb.active; ws.title = "Letters"
ws.append(["Letter", "Word_EN", "Word_ID", "Image", "Image_Prompt"])
for ch, en, idn in LETTERS:
    ws.append([ch, en, idn, "", f"cute {en.lower()}"])
style(ws, [8, 16, 18, 18, 30], "4FB3E8")

wt = wb.create_sheet("Templates")
wt.append(["Key", "Type", "Letters", "Title_EN", "Title_ID"])
for row in TEMPLATES:
    wt.append(list(row))
style(wt, [7, 10, 18, 56, 56], "9B7BE0")
dv = DataValidation(type="list", formula1='"capital,small,review,similar,pairs"', showErrorMessage=True,
                    error="Type: capital, small, review, similar, atau pairs")
wt.add_data_validation(dv); dv.add("B2:B200")

rd = wb.create_sheet("README")
for line in [
    "Sheet Letters: satu baris = satu huruf (A-Z) + kata kunci bergambar. Image (opsional) = nama file PNG tanpa .png; kosong = otomatis dari Word_EN (Apple -> apple.png, Ice Cream -> ice_cream.png).",
    "Gambar PNG ditaruh di assets/images (atau assets/Images). Belum ada PNG -> halaman memakai kotak teks biasa (demo).",
    "Sheet Templates: satu baris = satu halaman/produk. Key = huruf A-Z. Title_EN/Title_ID = judul halaman.",
    "Type: capital (huruf besar, 1-5 huruf) | small (huruf kecil, 1-5) | review (A-Z berurutan) | similar (huruf mirip) | pairs (Aa-Zz).",
    "Letters: rentang atau daftar dipisah koma: 'A-C' atau 'A-C, F'. Type similar memakai pasangan: 'b+d, p+q, m+n'.",
    "Kata di sheet Letters harus ada di data/approved_vocab.csv (terjemahan terverifikasi), kalau tidak QC memberi peringatan/gagal.",
]:
    rd.append([line])
rd.column_dimensions["A"].width = 170
wb.save("data/worksheet_database.xlsx"); print("tersimpan: data/worksheet_database.xlsx", len(LETTERS), "huruf,", len(TEMPLATES), "template")
